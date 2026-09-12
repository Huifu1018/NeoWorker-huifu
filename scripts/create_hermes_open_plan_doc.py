from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "neoworker-hermes-harness-open-plan.docx"

FONT = "FangSong"
TEXT = "202020"
MUTED = "5F6875"
ACCENT = "243B53"
LIGHT_BLUE = "EAF1F7"
LIGHT_GRAY = "F3F5F7"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is not None:
        tc_pr.remove(shd)
        tc_pr.insert(tc_pr.index(borders) + 1, shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            tc_pr.append(borders)
        else:
            tc_pr.insert(tc_pr.index(shd), borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        node = OxmlElement("w:keepNext")
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            p_pr.append(node)
        else:
            p_pr.insert(p_pr.index(spacing), node)
    elif not value and keep is not None:
        p_pr.remove(keep)


def set_keep_together(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepLines"))
    if value and keep is None:
        node = OxmlElement("w:keepLines")
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            p_pr.append(node)
        else:
            p_pr.insert(p_pr.index(spacing), node)
    elif not value and keep is not None:
        p_pr.remove(keep)


def set_run_font(run, name=FONT, size=None, bold=None, color=TEXT, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    lang.set(qn("w:bidi"), "ar-SA")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=6, line=1.18, left=0, first=0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.left_indent = Inches(left)
    fmt.first_line_indent = Inches(first)


def add_text(paragraph, text, *, bold=False, color=TEXT, size=10.5, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def add_body(doc, text, after=6):
    p = doc.add_paragraph(style="Normal")
    set_paragraph_format(p, after=after, line=1.2)
    add_text(p, text)
    set_keep_together(p)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, after=3, line=1.15)
    add_text(p, text)
    set_keep_together(p)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_format(p, after=3, line=1.15)
    add_text(p, text)
    set_keep_together(p)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    set_keep_with_next(p)
    set_keep_together(p)
    return p


def add_heading_text(doc, text, level=1):
    p = add_heading(doc, text, level)
    add_text(p, text, bold=True, color=TEXT, size=14 if level == 1 else 11.5)
    return p


def add_label_paragraph(doc, label, text, after=5):
    p = doc.add_paragraph(style="Normal")
    set_paragraph_format(p, after=after, line=1.18)
    add_text(p, label, bold=True)
    add_text(p, text)
    set_keep_together(p)
    return p


def add_table(doc, headers, rows, widths, font_size=9.5, alternate=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, (cell, value, width) in enumerate(zip(header.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, ACCENT)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_format(p, after=0, line=1.05)
        add_text(p, value, bold=True, color=WHITE, size=font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, value, width) in enumerate(zip(cells, values, widths)):
            set_cell_width(cell, width)
            set_cell_borders(cell)
            set_cell_margins(cell, top=105, start=130, bottom=105, end=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if alternate and row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, after=0, line=1.08)
            add_text(p, str(value), size=font_size)
            set_keep_together(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_small_note(doc, text):
    p = doc.add_paragraph(style="Normal")
    set_paragraph_format(p, before=0, after=5, line=1.12)
    add_text(p, text, color=MUTED, size=9)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title._element.rPr.rFonts.set(qn("w:cs"), FONT)
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(TEXT)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for level, size in ((1, 14), (2, 11.5), (3, 10.5)):
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(TEXT)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(5)
        style_ppr = style._element.get_or_add_pPr()
        num_pr = style_ppr.find(qn("w:numPr"))
        if num_pr is not None:
            style_ppr.remove(num_pr)
        border = style_ppr.find(qn("w:pBdr"))
        if border is not None:
            style_ppr.remove(border)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(TEXT)


def build_document():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_format(hp, after=0, line=1)
    add_text(hp, "NeoWorker Hermes Harness 开放计划", color=MUTED, size=8.5)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(fp, before=0, after=0, line=1)
    add_text(fp, "实施收口版  ·  2026年9月11日  ·  ", color=MUTED, size=8.5)
    add_page_field(fp)

    title = doc.add_paragraph(style="Title")
    add_text(title, "NeoWorker 使用 Hermes Harness 的开放计划", bold=True, color=TEXT, size=22)

    subtitle = doc.add_paragraph(style="Normal")
    set_paragraph_format(subtitle, after=12, line=1.1)
    add_text(subtitle, "技术策略草案 面向分阶段验证与评审", color=MUTED, size=11.5)

    metadata = [
        ["文档状态", "实施收口版，可根据后续评测和产品反馈调整"],
        ["日期", "2026年9月11日"],
        ["适用范围", "NeoWorker-huifu 的 Agent Runtime、工具执行和任务交付链"],
        ["核心建议", "新任务统一使用内置 Hermes ACP Harness；由 NeoWorker 保留工具、副作用和任务治理权"],
    ]
    add_table(doc, ["项目", "内容"], metadata, [1.12, 5.72], font_size=9.5, alternate=False)

    add_heading_text(doc, "一 结论", 1)
    add_body(
        doc,
        "NeoWorker 的新任务统一使用内置 Hermes ACP Harness。Hermes 负责 Agent Loop、规划和多步推进，NeoWorker 负责工具筛选、权限审批、沙箱、任务状态、证据记录和最终交付。Native SessionRuntime 仅作为历史任务兼容路径和明确的内部保底能力保留，不再作为新任务的按复杂度分流目标。",
    )
    add_body(
        doc,
        "这项改造已经落到 NeoWorker 的任务创建、IPC、运行时路由和执行器：新任务会绑定内置 Hermes ACP Host，前端不再暴露 Auto、Hermes、Native 运行时选择，Hermes 启动失败也不会静默切换到 Native。后续重点是持续扩大能力覆盖、完善真实任务评测和发布门禁，而不是再做一套面向用户的运行时选择器。",
    )
    add_label_paragraph(
        doc,
        "关键判断：",
        "“尽可能使用 Hermes”应理解为尽可能使用 Hermes 的 Agent Loop 和 Harness 能力，而不是把 Hermes 原生文件、Shell 和进程工具直接放回生产路径。生产任务的副作用仍应统一经过 NeoWorker Tool Host。",
    )

    add_heading_text(doc, "二 当前事实", 1)
    add_body(
        doc,
        "NeoWorker 需要区分“内置 Hermes Harness”和历史或显式外部 Runtime。它们的执行语义不同，不能只看 Provider 名称或 OpenAI-compatible 接口来判断是否真正使用了 Hermes Harness。",
    )
    add_table(
        doc,
        ["接入方式", "谁运行 Agent Loop", "谁执行本地副作用", "适合的定位"],
        [
            ["NeoWorker 内置 Hermes ACP Host", "Hermes", "NeoWorker Tool Host", "新任务默认路径，随安装包交付"],
            ["显式外部 ACP Runtime", "外部 Agent", "外部 Runtime 或 NeoWorker Host", "仅兼容旧任务和内部集成"],
            ["NeoWorker Native SessionRuntime", "NeoWorker", "NeoWorker", "历史任务兼容和内部保底路径"],
        ],
        [1.55, 1.25, 1.55, 2.49],
        font_size=9.1,
    )
    add_body(
        doc,
        "当前实现已经把新任务的默认路由固定到内置 Hermes：任务创建时绑定 Hermes ACP 配置，执行时通过随 NeoWorker 安装包交付的 Host 启动，不依赖用户另外安装 Hermes Agent。旧任务仍可能保留 Native 或外部 Runtime 元数据，因此历史记录中出现不同 Runtime 是兼容行为，不代表新任务还在按复杂度自动分流。",
    )
    add_body(
        doc,
        "专项验证已经覆盖内置 Hermes Host 的启动、ACP 工具桥接、审批、工具生命周期、checkpoint、恢复、Shell 和安装包 smoke。当前状态可以概括为：新任务统一使用内置 Hermes，Native 仅保留兼容语义，后续工作集中在真实任务质量和能力覆盖。",
    )
    add_small_note(
        doc,
        "依据：docs/hermes-runtime.md、docs/hermes-tool-ownership-audit.md 和 docs/hermes-test-report.md。上述文件记录了接入方式、工具所有权以及专项验证范围。",
    )

    add_heading_text(doc, "三 目标状态", 1)
    add_body(
        doc,
        "目标状态不是让两个 Runtime 互相覆盖，而是让它们形成清晰的分工。NeoWorker 负责“任务和工具的边界”，Hermes 负责“如何把复杂任务推进到完成”。",
    )
    add_table(
        doc,
        ["能力层", "建议归属", "目标行为"],
        [
            ["Agent Loop 和任务规划", "Hermes ACP", "Hermes 负责拆解目标、选择下一步、持续推进和处理工具结果"],
            ["工具目录和任务级筛选", "NeoWorker", "只向 Hermes 暴露当前任务获准使用的工具"],
            ["文件、Shell、Office 和其他副作用", "NeoWorker", "所有执行经过现有 Tool Host、审批、沙箱、超时和日志链"],
            ["任务状态和恢复", "NeoWorker 主导，Hermes 提供 session", "任务可暂停、恢复、重试，未知副作用必须显式确认"],
            ["用户界面和运行时状态", "NeoWorker", "不展示运行时选择；时间线只展示实际运行状态、失败和历史回退"],
            ["最终交付校验", "NeoWorker", "文件类任务只有在产物存在、可解析和通过基础检查后才能完成"],
        ],
        [1.65, 1.55, 3.64],
        font_size=9.1,
    )

    add_heading_text(doc, "四 运行时策略", 1)
    add_body(
        doc,
        "运行时选择收回后端。新任务不再由前端选择，也不再按任务复杂度在 Hermes 和 Native 之间自适应切换；后端统一创建内置 Hermes 任务。历史任务和显式外部集成继续按已持久化的 Runtime 运行，并保留兼容性。",
    )
    add_table(
        doc,
        ["任务来源", "运行时", "适用情况", "失败处理"],
        [
            ["新任务", "NeoWorker 内置 Hermes ACP", "所有普通对话、网页、代码和 Office 任务", "Hermes 启动或执行失败即失败，不静默切 Native"],
            ["历史 Native 任务", "NeoWorker Native SessionRuntime", "保持已有会话和历史任务语义", "沿用原生暂停、恢复和错误语义"],
            ["显式外部 ACP 任务", "按外部 Runtime 配置", "旧集成、内部调度或明确委托的任务", "不改变外部契约；禁止无记录的隐式替换"],
        ],
        [1.35, 1.72, 2.32, 1.45],
        font_size=9.0,
    )
    add_label_paragraph(
        doc,
        "路由原则：",
        "新任务的唯一默认判断是“使用内置 Hermes Harness”；任务复杂度只影响 Hermes 内部的规划和工具调用，不再决定是否切换到 Native。运行时事件必须记录实际 Runtime、状态和失败原因。",
    )

    add_heading_text(doc, "五 任务路由建议", 1)
    add_table(
        doc,
        ["任务类型", "建议 Runtime", "原因和边界"],
        [
            ["网页查询和结构化数据采集", "内置 Hermes ACP", "需要搜索、抓取、解析、去重、校验和总结等连续步骤"],
            ["代码、终端和调试", "内置 Hermes ACP", "由 Hermes 规划和试错；命令仍由 NeoWorker 执行"],
            ["PPT、Word 和其他 Office 多步任务", "内置 Hermes ACP", "先理解模板、再修改、再检查；必须增加产物完成门禁"],
            ["简单问答和短文本改写", "内置 Hermes ACP", "保持统一任务语义，不再为省启动开销切换 Native"],
            ["单次确定性工具调用", "内置 Hermes ACP", "仍由统一 Harness 承接，便于审计、会话和结果收口"],
            ["依赖 Hermes 原生 Memory、Project Plugin 或特殊 Toolset 的任务", "内置 Hermes ACP + 能力评估", "当前宿主包装器保持受控能力范围，需先完成能力映射和所有权评估"],
        ],
        [1.9, 1.15, 3.79],
        font_size=8.9,
    )
    add_body(
        doc,
        "PPT 截图暴露出的“工具调用已经结束，但回答仍停在准备读取模板”属于另一个需要单独修复的问题：它可能是任务终态、产物校验或输出可见性问题，不能简单归因于是否使用 Hermes。无论 Runtime 选择什么，Office 任务都应在标记完成前验证最终文件。",
    )

    add_heading_text(doc, "六 分阶段开放计划", 1)
    add_body(
        doc,
        "下面的阶段按风险从低到高排列。每一阶段都有明确的产出和退出条件；如果评测结果不理想，可以停留在当前阶段，不必一次性做全量切换。",
    )

    add_heading_text(doc, "阶段 0 建立可见性和运行时契约", 2)
    add_body(doc, "先让用户和开发者能准确知道每个任务使用了什么 Runtime，以及任务为什么选择这条路径。")
    add_bullet(doc, "新任务由后端统一绑定内置 Hermes ACP 配置；前端不再展示或提交 Auto、Hermes、Native 运行时选择。")
    add_bullet(doc, "时间线显示实际 Hermes、历史 Native、失败和历史回退状态；记录 external runtime、Hermes 版本和宿主工具模式。")
    add_bullet(doc, "内置 Hermes 启动失败或执行失败时直接结束为失败，不静默切换到 Native；错误信息明确说明无需额外安装 Hermes Agent。")
    add_bullet(doc, "建立统一的完成契约，特别是 Office 文件任务必须通过文件存在、大小、解析和必要的预览检查。")
    add_label_paragraph(doc, "退出条件：", "用户可以从任务记录中确认实际 Runtime；新任务不会被误标为 Native；内置 Hermes 的失败和历史回退可追溯。")

    add_heading_text(doc, "阶段 1 固定新任务进入内置 Hermes", 2)
    add_body(doc, "这一步已经完成：所有普通新任务统一使用随 NeoWorker 交付的 Hermes ACP Host，不需要额外部署 Hermes Agent。")
    add_bullet(doc, "任务创建、Renderer IPC 和 AgentDaemon 创建路径统一写入内置 Hermes external runtime。")
    add_bullet(doc, "继续复用现有 Tool Host、审批、沙箱、checkpoint、暂停恢复、取消和日志链。")
    add_bullet(doc, "保留 Native 仅用于历史任务兼容和明确的内部保底语义，不作为新任务的自动降级目标。")
    add_bullet(doc, "为 Hermes 任务设置并发、超时、取消和资源上限；继续复用现有 Tool Host、审批、沙箱和日志。")
    add_bullet(doc, "收集真实任务样本，优先覆盖航班查询、网页抓取、代码修改、PPT 模板填充和文件整理。")
    add_label_paragraph(doc, "退出条件：", "候选任务的 Hermes 成功率、完成质量、耗时和用户修正成本不劣于 Native；没有未记录的本地副作用。具体数值在阶段 0 建立基线后确定。")

    add_heading_text(doc, "阶段 2 补齐 Hermes 能力覆盖", 2)
    add_body(doc, "当前包装器为了保证工具所有权，会关闭 Hermes workspace context、memory、project plugins 和 inherited kanban。阶段 2 要逐项判断哪些能力值得恢复，哪些能力应该由 NeoWorker 提供等价实现。")
    add_bullet(doc, "建立 Hermes 能力清单，区分 Agent Loop 能力、上下文能力、Memory 能力、插件能力和原生副作用工具。")
    add_bullet(doc, "优先把可治理的上下文和记忆能力接入 NeoWorker，而不是直接开放 Hermes 的本地副作用工具。")
    add_bullet(doc, "为每项新增能力写明数据来源、权限边界、任务范围、审计字段和恢复语义。")
    add_bullet(doc, "保留版本固定和 fail-closed 策略，在 Hermes 版本升级后重新运行 ACP、宿主工具、Windows Shell 和安装包 smoke。")
    add_label_paragraph(doc, "退出条件：", "需要 Hermes 能力的目标任务可以通过 NeoWorker 公开的治理接口获得相同或足够接近的结果，且所有权和恢复边界仍然清楚。")

    add_heading_text(doc, "阶段 3 完善能力覆盖并优化体验", 2)
    add_body(doc, "在统一 Harness 路径稳定后，继续补齐 Hermes 能力映射、结果收口和用户可理解性。")
    add_bullet(doc, "根据真实任务评测扩展 Hermes 的受控工具、上下文和 Office 产物校验能力。")
    add_bullet(doc, "继续利用 ACP 热会话、checkpoint 和受控重试，减少连续任务的启动开销。")
    add_bullet(doc, "把 Hermes 运行状态、工具审批、降级和产物检查统一成同一套任务视图。")
    add_bullet(doc, "将 Hermes 版本、模型、宿主工具协议和安装包 smoke 纳入固定发布门禁。")
    add_label_paragraph(doc, "退出条件：", "新任务稳定经过内置 Hermes Harness，Native 兼容路径不影响新任务，用户能在时间线中看懂实际状态和失败原因。")

    add_heading_text(doc, "七 评测和发布门槛", 1)
    add_body(
        doc,
        "这项改造不能只用“能否启动 Hermes”来验收。核心问题是：任务是否完成得更好，副作用是否仍然受 NeoWorker 控制，用户是否能看懂运行时状态。",
    )
    add_table(
        doc,
        ["指标", "需要观察的结果", "最低门槛"],
        [
            ["Runtime 真实性", "任务实际进入的 Runtime 与界面和事件记录一致", "新任务无法被静默替换为 Native"],
            ["工具所有权", "文件、Shell、Office 等副作用都能关联到 NeoWorker Tool Host", "零未知工具调用，零未记录副作用"],
            ["任务完成质量", "结果完整、准确、结构清楚，能通过用户验收", "不低于 Native 基线；复杂任务单独比较"],
            ["产物可靠性", "生成的 DOCX、PPTX 等文件存在且可解析", "未通过产物检查不得标记完成"],
            ["恢复安全", "暂停、取消、重试和重启不会自动重放未知副作用", "需要显式确认，且 checkpoint 可追踪"],
            ["体验和效率", "首字节、总耗时、工具审批次数和模型成本可比较", "建立基线后设定可接受范围"],
        ],
        [1.42, 3.75, 1.71],
        font_size=8.9,
    )
    add_label_paragraph(
        doc,
        "建议评测集：",
        "至少包含网页结构化采集、航班或商品查询、代码修改与测试、PPT 模板填充、Word 内容整理、文件批处理和需要审批的 Shell 操作。每类任务都应有 Native 与 Hermes ACP 的可比样本。",
    )

    add_heading_text(doc, "八 主要风险", 1)
    add_table(
        doc,
        ["风险", "表现", "应对"],
        [
            ["能力不完全等价", "当前包装器关闭 Hermes 的部分上下文、Memory 和插件能力", "先建立能力清单，按价值和治理成本逐项迁移"],
            ["运行时误判", "用户以为用了 Hermes，实际任务走了 Native 或静默降级", "运行时持久化，时间线显示，新任务禁止静默降级"],
            ["工具所有权混乱", "8642 或原生 ACP toolset 直接执行 Hermes 工具", "生产默认使用 ACP 宿主工具模式，副作用统一进入 NeoWorker Tool Host"],
            ["版本耦合", "当前要求 hermes-agent 0.18.0，升级可能改变 ACP 合同", "版本固定、fail-closed、升级后专项回归和安装包 smoke"],
            ["结果质量误归因", "模型、Prompt、数据源或工具差异被误认为 Harness 差异", "A/B 评测尽量固定模型、输入、工具目录和数据条件"],
            ["Office 任务终态不可靠", "工具调用完成但文件未生成或最终回答未收口", "增加产物校验和统一终态收口，Runtime 无关地执行"],
        ],
        [1.35, 3.06, 2.47],
        font_size=8.75,
    )

    add_heading_text(doc, "九 待确认事项", 1)
    add_number(doc, "内置 Hermes Host 在不同机器和安装包形态下的启动成功率、耗时和资源占用是否满足发布门槛？")
    add_number(doc, "Hermes 的 Memory、workspace context 和 project plugins 中，哪些能力是必须保留的？哪些可以由 NeoWorker 提供等价接口？")
    add_number(doc, "历史 Native 任务和显式外部 ACP 任务的兼容边界是否需要进一步收紧？")
    add_number(doc, "模型和 Provider 是否在 A/B 评测中固定？如果不固定，需要把模型差异单独记录，避免把模型效果算到 Harness 上。")
    add_number(doc, "是否把内置 Hermes Host 版本升级设为独立发布门禁？每次升级都必须重新验收 ACP、工具、恢复和安装包行为。")

    add_heading_text(doc, "十 建议的第一批工作", 1)
    add_body(
        doc,
        "第一批工作已经完成运行时统一和失败语义收口。后续工作不扩大副作用边界，重点是把内置 Hermes 的能力覆盖、结果质量和发布验证做扎实。",
    )
    add_bullet(doc, "新任务默认绑定内置 Hermes，前端隐藏运行时选择，后端保留旧字段仅用于兼容历史数据。")
    add_bullet(doc, "增加 Hermes Active、历史 Native、Failed 和历史 Fallback 的时间线事件及 UI 标识。")
    add_bullet(doc, "为内置 Hermes 任务关闭静默 fallback；运行时不可用时明确失败并提示无需单独安装 Hermes Agent。")
    add_bullet(doc, "建立第一版 A/B 评测集，并把航班查询和 PPT 模板任务纳入产物校验。")
    add_bullet(doc, "在本地和 CI 中继续保留 ACP、Tool Host、审批、Shell、checkpoint、Windows 和安装包验证。")
    add_label_paragraph(
        doc,
        "第一阶段的判断标准：",
        "用户能够明确知道任务是否真正使用 Hermes，所有新任务能够稳定经过 Hermes Agent Loop，NeoWorker 仍然拥有所有工具副作用，并且任何失败或历史回退都不会被伪装成成功。",
    )

    add_heading_text(doc, "附录 依据文件", 1)
    add_small_note(doc, "docs/hermes-runtime.md 记录内置 Hermes ACP Host、Tool Host 边界和运行时生命周期。")
    add_small_note(doc, "docs/hermes-tool-ownership-audit.md 记录 Hermes Agent、Hermes Model Proxy 和 Hermes ACP 的工具所有权差异。")
    add_small_note(doc, "docs/hermes-test-report.md 记录 ACP 宿主工具、审批、checkpoint、恢复、Shell、跨平台和安装包专项验证。")
    add_small_note(doc, "src/electron/agent/executor.ts 包含内置 Hermes external runtime 的执行分支；Native SessionRuntime 仅保留历史兼容路径。")

    # Avoid widows around the final section and keep the document clean when opened in Word.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            set_keep_with_next(paragraph)
        for run in paragraph.runs:
            set_run_font(run, name=FONT, size=run.font.size.pt if run.font.size else None)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
